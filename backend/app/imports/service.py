from __future__ import annotations

import csv
import json
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SECTIONS = ("experience", "education", "skills", "projects")
LIST_FIELDS = {"highlights", "keywords", "technologies"}


def _clean(value: Any) -> Any:
    if value is None:
        return ""
    return value.strip() if isinstance(value, str) else value


def _rows_from_xlsx(path: Path) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: dict[str, list[dict[str, Any]]] = {}
    try:
        for section in SECTIONS:
            sheet_name = section.title()
            if sheet_name not in workbook.sheetnames:
                rows[section] = []
                continue
            values = list(workbook[sheet_name].iter_rows(values_only=True))
            if not values:
                rows[section] = []
                continue
            headers = [str(value or "").strip().lower() for value in values[0]]
            rows[section] = [
                {headers[index]: _clean(value) for index, value in enumerate(row) if headers[index]}
                for row in values[1:]
                if any(value not in (None, "") for value in row)
            ]
    finally:
        workbook.close()
    return rows


def _rows_from_csv(path: Path, section: str) -> dict[str, list[dict[str, Any]]]:
    if section not in SECTIONS:
        raise ValueError(f"CSV imports require --section: {', '.join(SECTIONS)}")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {name: list(csv.DictReader(handle)) if name == section else [] for name in SECTIONS}


def read_import(path: Path, section: str | None = None) -> dict[str, list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _rows_from_xlsx(path)
    if suffix == ".csv":
        return _rows_from_csv(path, section or "")
    raise ValueError("Import file must be .xlsx or .csv")


def _normalize(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {key: _clean(value) for key, value in row.items()}
    for field in LIST_FIELDS:
        if field in normalized and isinstance(normalized[field], str):
            normalized[field] = [part.strip() for part in normalized[field].split("|") if part.strip()]
    return normalized


def build_review(source: Path, canonical: Path, section: str | None = None) -> dict[str, Any]:
    data = json.loads(canonical.read_text(encoding="utf-8"))
    imported = read_import(source, section)
    changes: list[dict[str, Any]] = []
    errors: list[str] = []
    for section_name, rows in imported.items():
        existing = {item.get("id"): item for item in data.get(section_name, []) if item.get("id")}
        for row_number, raw in enumerate(rows, start=2):
            row = _normalize(raw)
            record_id = row.pop("id", "")
            operation = str(row.pop("operation", "upsert") or "upsert").lower()
            source_note = row.pop("source", "spreadsheet") or "spreadsheet"
            note = row.pop("note", "")
            if not record_id:
                errors.append(f"{section_name} row {row_number}: id is required")
                continue
            if operation not in {"upsert", "archive"}:
                errors.append(f"{section_name} row {row_number}: operation must be upsert or archive")
                continue
            before = existing.get(record_id)
            if operation == "archive" and before is None:
                errors.append(f"{section_name} row {row_number}: cannot archive unknown id {record_id}")
                continue
            after = deepcopy(before) if before else {"id": record_id}
            if operation == "archive":
                after["status"] = "archived"
            else:
                after.update({key: value for key, value in row.items() if value not in ("", [])})
                after["status"] = row.get("status") or after.get("status") or "active"
            comparable_before = {key: value for key, value in (before or {}).items() if key != "provenance"}
            comparable_after = {key: value for key, value in after.items() if key != "provenance"}
            if before is not None and comparable_before == comparable_after:
                continue
            after["provenance"] = {
                "source": str(source_note),
                "source_type": "spreadsheet" if source.suffix.lower() == ".xlsx" else "import",
                "imported_at": datetime.now(timezone.utc).isoformat(),
                **({"note": str(note)} if note else {}),
            }
            result = "archive" if operation == "archive" else ("update" if before else "create")
            if before != after:
                changes.append({"section": section_name, "id": record_id, "result": result, "before": before, "after": after})
    return {
        "source": str(source),
        "canonical": str(canonical),
        "valid": not errors,
        "summary": {
            "create": sum(change["result"] == "create" for change in changes),
            "update": sum(change["result"] == "update" for change in changes),
            "archive": sum(change["result"] == "archive" for change in changes),
            "errors": len(errors),
        },
        "errors": errors,
        "changes": changes,
    }


def apply_import(review: dict[str, Any], canonical: Path) -> dict[str, Any]:
    if not review["valid"]:
        raise ValueError("Cannot apply an import with validation errors")
    data = json.loads(canonical.read_text(encoding="utf-8"))
    for change in review["changes"]:
        records = data.setdefault(change["section"], [])
        for index, record in enumerate(records):
            if record.get("id") == change["id"]:
                records[index] = change["after"]
                break
        else:
            records.append(change["after"])
    data.setdefault("schema_version", "1.0")
    canonical.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return data
