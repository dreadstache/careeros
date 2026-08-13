import json
import subprocess
from pathlib import Path

import pytest
from openpyxl import Workbook

import app.main as main_module
import app.studio as studio_module
from app.imports import apply_import, build_review
from app.main import app
from app.studio import export_ecosystem_manifest, export_track_manifest, publishing_status, save_tracks
from fastapi.testclient import TestClient


def _canonical(path: Path) -> Path:
    path.write_text(json.dumps({"experience": [{"id": "old-role", "status": "active", "organization": "Old Co"}]}), encoding="utf-8")
    return path


def test_review_does_not_mutate_and_reports_create_update_archive(tmp_path):
    canonical = _canonical(tmp_path / "career.json")
    before = canonical.read_text(encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Experience"
    sheet.append(["id", "operation", "status", "organization", "position", "source", "note"])
    sheet.append(["old-role", "archive", "active", "Old Co", "", "career workbook", "retired"])
    sheet.append(["new-role", "upsert", "active", "New Co", "Artist", "career workbook", "verified"])
    source = tmp_path / "import.xlsx"
    workbook.save(source)

    review = build_review(source, canonical)

    assert review["valid"] is True
    assert review["summary"] == {"create": 1, "update": 0, "archive": 1, "errors": 0}
    assert canonical.read_text(encoding="utf-8") == before

    applied = apply_import(review, canonical)
    assert {item["id"]: item["status"] for item in applied["experience"]} == {"old-role": "archived", "new-role": "active"}


def test_review_rejects_missing_id(tmp_path):
    canonical = _canonical(tmp_path / "career.json")
    source = tmp_path / "skills.csv"
    source.write_text("id,operation,name\n,upsert,Python\n", encoding="utf-8")
    review = build_review(source, canonical, "skills")
    assert review["valid"] is False
    assert review["summary"]["errors"] == 1


def test_unchanged_row_is_not_reported(tmp_path):
    canonical = _canonical(tmp_path / "career.json")
    source = tmp_path / "experience.csv"
    source.write_text("id,operation,status,organization\nold-role,upsert,active,Old Co\n", encoding="utf-8")
    review = build_review(source, canonical, "experience")
    assert review["summary"] == {"create": 0, "update": 0, "archive": 0, "errors": 0}


def test_review_endpoint_accepts_workbook_without_mutating(tmp_path, monkeypatch):
    canonical = _canonical(tmp_path / "career.json")
    monkeypatch.setattr("app.main.CANONICAL", canonical)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Experience"
    sheet.append(["id", "operation", "status", "organization"])
    sheet.append(["new-role", "upsert", "active", "New Co"])
    source = tmp_path / "import.xlsx"
    workbook.save(source)
    before = canonical.read_text(encoding="utf-8")
    response = TestClient(app).post("/imports/review", files={"file": ("career.xlsx", source.read_bytes())})
    assert response.status_code == 200
    assert response.json()["summary"]["create"] == 1
    assert canonical.read_text(encoding="utf-8") == before


def test_excel_numeric_years_are_normalized_to_strings(tmp_path):
    canonical = _canonical(tmp_path / "career.json")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Experience"
    sheet.append(["id", "operation", "status", "organization", "start_date", "end_date"])
    sheet.append(["new-role", "upsert", "active", "New Co", 2024, 2026])
    source = tmp_path / "import.xlsx"
    workbook.save(source)
    review = build_review(source, canonical)
    assert review["changes"][0]["after"]["start_date"] == "2024"
    assert review["changes"][0]["after"]["end_date"] == "2026"


def _studio_root(tmp_path: Path) -> Path:
    (tmp_path / "data").mkdir()
    (tmp_path / "frontend" / "public" / "generated").mkdir(parents=True)
    (tmp_path / "data" / "career-data.json").write_text(
        json.dumps({"experience": [{"id": "old-role", "status": "active", "organization": "Old Co"}]}),
        encoding="utf-8",
    )
    (tmp_path / "forge.resume.json").write_text(
        json.dumps({
            "project_name": "generated",
            "module_options": {"resume": {"profiles": [{
                "slug": "example",
                "title": "Example",
                "headline": "Example headline",
                "summary": "Example summary",
                "experience_ids": ["old-role"],
                "skill_ids": [],
                "project_ids": [],
            }]}},
        }),
        encoding="utf-8",
    )
    (tmp_path / "data" / "ecosystem.json").write_text(
        json.dumps({
            "schema_version": "1.0",
            "identity": {"name": "Example Person"},
            "destinations": [{
                "id": "portfolio",
                "label": "Portfolio",
                "description": "Example destination",
                "url": "https://example.com/",
                "status": "live",
            }],
        }),
        encoding="utf-8",
    )
    return tmp_path


def _review_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Experience"
    sheet.append(["id", "operation", "status", "organization", "position"])
    sheet.append(["new-role", "upsert", "active", "New Co", "Developer"])
    from io import BytesIO
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_studio_apply_requires_owner_token_and_regenerates(tmp_path, monkeypatch):
    root = _studio_root(tmp_path)
    monkeypatch.setattr(main_module, "ROOT", root)
    monkeypatch.setattr(main_module, "CANONICAL", root / "data" / "career-data.json")
    monkeypatch.setattr(studio_module, "run_generation", lambda _root: "Generated modules: base, resume")
    monkeypatch.setenv("CAREEROS_OWNER_TOKEN", "test-owner-token")
    studio_module.PENDING_REVIEWS.clear()
    client = TestClient(app)

    review = client.post("/imports/review", files={"file": ("career.xlsx", _review_workbook())})
    assert review.status_code == 200
    review_id = review.json()["review_id"]
    assert client.post(f"/imports/{review_id}/apply").status_code == 401

    applied = client.post(
        f"/imports/{review_id}/apply",
        headers={"X-CareerOS-Owner-Token": "test-owner-token"},
    )
    assert applied.status_code == 200
    assert applied.json()["status"] == "applied"
    canonical = json.loads((root / "data" / "career-data.json").read_text(encoding="utf-8"))
    assert {item["id"] for item in canonical["experience"]} == {"old-role", "new-role"}
    assert list((root / "backups").glob("career-data-*.json"))


def test_studio_rejects_review_when_canonical_changed(tmp_path, monkeypatch):
    root = _studio_root(tmp_path)
    monkeypatch.setattr(main_module, "ROOT", root)
    monkeypatch.setattr(main_module, "CANONICAL", root / "data" / "career-data.json")
    monkeypatch.setenv("CAREEROS_OWNER_TOKEN", "test-owner-token")
    studio_module.PENDING_REVIEWS.clear()
    client = TestClient(app)
    review = client.post("/imports/review", files={"file": ("career.xlsx", _review_workbook())}).json()
    (root / "data" / "career-data.json").write_text(json.dumps({"experience": []}), encoding="utf-8")

    response = client.post(
        f"/imports/{review['review_id']}/apply",
        headers={"X-CareerOS-Owner-Token": "test-owner-token"},
    )
    assert response.status_code == 409
    assert "Review the workbook again" in response.json()["detail"]


def test_track_selections_validate_ids_and_regenerate(tmp_path, monkeypatch):
    root = _studio_root(tmp_path)
    monkeypatch.setattr(studio_module, "run_generation", lambda _root: "generated")
    track = {
        "slug": "example",
        "original_slug": "example",
        "title": "Example",
        "headline": "Example headline",
        "summary": "Example summary",
        "experience_ids": ["old-role"],
        "skill_ids": [],
        "project_ids": [],
    }
    result = save_tracks(root, [track], [], False)
    assert result["status"] == "saved"
    assert result["tracks"][0]["experience_ids"] == ["old-role"]
    with pytest.raises(ValueError, match="must include at least one"):
        save_tracks(root, [{**track, "experience_ids": []}], [], False)
    with pytest.raises(ValueError, match="unknown experience_ids"):
        save_tracks(root, [{**track, "experience_ids": ["invented-role"]}], [], False)


def test_tracks_can_be_added_renamed_and_reordered(tmp_path, monkeypatch):
    root = _studio_root(tmp_path)
    monkeypatch.setattr(studio_module, "run_generation", lambda _root: "generated")
    tracks = [
        {
            "slug": "new-focus",
            "original_slug": None,
            "title": "New Focus Résumé",
            "headline": "A new focus",
            "summary": "A focused summary",
            "experience_ids": ["old-role"],
            "skill_ids": [],
            "project_ids": [],
        },
        {
            "slug": "renamed-example",
            "original_slug": "example",
            "title": "Renamed Example Résumé",
            "headline": "Renamed headline",
            "summary": "Renamed summary",
            "experience_ids": ["old-role"],
            "skill_ids": [],
            "project_ids": [],
        },
    ]

    result = save_tracks(root, tracks, [], False)

    assert [track["slug"] for track in result["tracks"]] == ["new-focus", "renamed-example"]
    manifest = json.loads((root / "frontend" / "public" / "generated" / "resume" / "tracks.json").read_text(encoding="utf-8"))
    assert [track["slug"] for track in manifest["tracks"]] == ["new-focus", "renamed-example"]


def test_track_removal_requires_exact_confirmation_and_removes_stale_output(tmp_path, monkeypatch):
    root = _studio_root(tmp_path)
    stale = root / "frontend" / "public" / "generated" / "resume" / "example"
    stale.mkdir(parents=True)
    (stale / "index.html").write_text("old", encoding="utf-8")
    monkeypatch.setattr(studio_module, "run_generation", lambda _root: "generated")
    replacement = {
        "slug": "replacement",
        "original_slug": None,
        "title": "Replacement Résumé",
        "headline": "Replacement headline",
        "summary": "Replacement summary",
        "experience_ids": ["old-role"],
        "skill_ids": [],
        "project_ids": [],
    }

    with pytest.raises(ValueError, match="Confirm track removal"):
        save_tracks(root, [replacement], ["example"], False)
    with pytest.raises(ValueError, match="does not match"):
        save_tracks(root, [replacement], [], True)

    save_tracks(root, [replacement], ["example"], True)
    assert not stale.exists()


def test_export_track_manifest_keeps_profile_order(tmp_path):
    root = _studio_root(tmp_path)
    destination = export_track_manifest(root)
    manifest = json.loads(destination.read_text(encoding="utf-8"))
    assert manifest == {
        "schema_version": "1.0",
        "tracks": [{
            "slug": "example",
            "title": "Example",
            "headline": "Example headline",
            "summary": "Example summary",
        }],
    }


def test_export_ecosystem_manifest_validates_and_copies_public_navigation(tmp_path):
    root = _studio_root(tmp_path)
    destination = export_ecosystem_manifest(root)
    manifest = json.loads(destination.read_text(encoding="utf-8"))
    assert manifest["identity"]["name"] == "Example Person"
    assert manifest["destinations"][0]["url"] == "https://example.com/"

    source = root / "data" / "ecosystem.json"
    invalid = json.loads(source.read_text(encoding="utf-8"))
    invalid["destinations"][0]["url"] = "http://example.com/"
    source.write_text(json.dumps(invalid), encoding="utf-8")
    with pytest.raises(ValueError, match="requires an HTTPS URL"):
        export_ecosystem_manifest(root)


def test_publishing_status_allows_only_career_data_and_config(tmp_path):
    root = _studio_root(tmp_path)
    subprocess.run(["git", "init", "-b", "main"], cwd=root, check=True, capture_output=True)
    subprocess.run(["git", "config", "user.email", "studio@example.com"], cwd=root, check=True)
    subprocess.run(["git", "config", "user.name", "CareerOS Studio"], cwd=root, check=True)
    subprocess.run(["git", "add", "."], cwd=root, check=True)
    subprocess.run(["git", "commit", "-m", "fixture"], cwd=root, check=True, capture_output=True)
    canonical = root / "data" / "career-data.json"
    canonical.write_text(canonical.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    status = publishing_status(root)
    assert status["ready"] is True
    assert status["publishable_changes"] == ["data/career-data.json"]

    (root / "README.md").write_text("unrelated", encoding="utf-8")
    status = publishing_status(root)
    assert status["ready"] is False
    assert status["unrelated_changes"] == ["README.md"]
